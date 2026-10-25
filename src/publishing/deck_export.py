"""Exports a consolidated course as a .pptx against docs/contracts/deck-template.md.

The code names no template, layout or placeholder index; those come from the mapping file.
"""
import io
import os
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, List, Optional

import yaml
from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER

from src.publishing.image_ids import ImageRefError
from src.publishing.course import (
    Course, DeckExportError, course_concepts, decks, format_citation, format_deck_line,
    iter_leaves, iter_slides, notes_text, title_body,
)

DECK_TEMPLATE_VERSION = 1
NOTES_PAGE = "notes-page"

DEFAULT_MAPPING = Path(
    os.getenv("DECK_TEMPLATE")
    or Path(__file__).resolve().parents[2] / "tests" / "fixtures" / "templates" / "standin" / "deck-template.yaml"
)

ROLES = {
    "title": {"required": ["title"], "optional": ["body"]},
    "section": {"required": ["title"], "optional": ["body"]},
    "concept": {"required": ["title", "body", "citation"], "optional": ["notes"]},
    "concept-with-figure": {"required": ["title", "body", "citation", "image"], "optional": ["notes"]},
    "summary": {"required": ["title", "body"], "optional": []},
    "sources": {"required": ["title", "body"], "optional": []},
}
IMAGE_TYPES = (PP_PLACEHOLDER.PICTURE, PP_PLACEHOLDER.OBJECT)


@dataclass
class Mapping:
    path: Path
    version: object
    template: Optional[str]
    roles: Dict[str, dict]

    @property
    def template_path(self) -> Optional[Path]:
        return (self.path.parent / self.template) if self.template else None


def load_mapping(path) -> Mapping:
    path = Path(path)
    try:
        raw = yaml.safe_load(path.read_text(encoding="utf-8"))
    except (OSError, yaml.YAMLError) as e:
        raise DeckExportError([f"mapping file {path} could not be read: {e}"])
    if not isinstance(raw, dict):
        raise DeckExportError([f"mapping file {path} is not a mapping"])
    return Mapping(path, raw.get("deck_template"), raw.get("template"), raw.get("roles") or {})


def _layouts_by_name(prs):
    out: Dict[str, list] = {}
    for master in prs.slide_masters:
        for layout in master.slide_layouts:
            out.setdefault(layout.name, []).append(layout)
    return out


def _needed_roles(course):
    return list(dict.fromkeys(role for role, _ in iter_slides(course)))


def validate(mapping: Mapping, course) -> List[str]:
    problems: List[str] = []
    if mapping.version != DECK_TEMPLATE_VERSION:
        problems.append(f"deck_template is {mapping.version!r}, this exporter reads {DECK_TEMPLATE_VERSION}")
    prs = None
    if not mapping.template:
        problems.append("the mapping names no template")
    elif not mapping.template_path.is_file():
        problems.append(f"template file {mapping.template_path} is missing")
    else:
        try:
            prs = Presentation(str(mapping.template_path))
        except Exception as e:
            problems.append(f"template file {mapping.template_path} could not be opened: {e}")

    for role in _needed_roles(course):
        if role not in mapping.roles:
            problems.append(f"role '{role}' is not mapped")

    layouts = _layouts_by_name(prs) if prs else {}
    for role, spec in mapping.roles.items():
        if role not in ROLES:
            problems.append(f"unknown role '{role}'")
            continue
        spec = spec or {}
        fields = spec.get("placeholders") or {}
        known = ROLES[role]["required"] + ROLES[role]["optional"]
        for f in fields:
            if f not in known:
                problems.append(f"role '{role}': unknown field '{f}'")
        for f in ROLES[role]["required"]:
            if f not in fields:
                problems.append(f"role '{role}': required field '{f}' is not mapped")
        if role == "title" and "body" not in fields:
            problems.append(f"role 'title': the course has a subtitle and the role maps no 'body' field for it")

        name = spec.get("layout")
        if not name:
            problems.append(f"role '{role}': no layout named")
            continue
        if not prs:
            continue
        found = layouts.get(name, [])
        if not found:
            problems.append(f"role '{role}': layout '{name}' is not in the template; its layouts are: "
                            + ", ".join(sorted(layouts)))
            continue
        if len(found) > 1:
            problems.append(f"role '{role}': layout '{name}' is carried by {len(found)} layouts in the template")
            continue
        types = {p.placeholder_format.idx: p.placeholder_format.type for p in found[0].placeholders}
        used: Dict[int, str] = {}
        for f, idx in fields.items():
            if f not in known:
                continue
            if f == "notes" and idx == NOTES_PAGE:
                continue
            if isinstance(idx, bool) or not isinstance(idx, int):
                problems.append(f"role '{role}': field '{f}' must be a placeholder idx, got {idx!r}")
                continue
            if idx not in types:
                problems.append(f"role '{role}': field '{f}' names placeholder {idx}, which layout '{name}' does not have")
                continue
            if f == "image" and types[idx] not in IMAGE_TYPES:
                problems.append(f"role '{role}': image placeholder {idx} on layout '{name}' cannot take a picture")
            if idx in used:
                problems.append(f"role '{role}': placeholder {idx} is mapped to both '{used[idx]}' and '{f}'")
            used[idx] = f
    return problems


def _runs(text):
    parts = re.split(r"(\*\*.+?\*\*)", text)
    for part in parts:
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            yield part[2:-2], True
        elif part:
            yield part, False


def write_text(text_frame, markdown):
    """Paragraphs from markdown lines: bullets and indentation become levels, `**x**` and `#` headers become bold."""
    lines = [ln.rstrip() for ln in (markdown or "").splitlines() if ln.strip()]
    text_frame.clear()
    first = True
    for line in lines:
        indent = len(line.replace("\t", "    ")) - len(line.replace("\t", "    ").lstrip())
        content = line.strip()
        bold_all = False
        m = re.match(r"^(#+)\s+(.*)$", content)
        if m:
            content, bold_all = m.group(2), True
            indent = 0
        else:
            m = re.match(r"^[-*+]\s+(.*)$", content)
            if m:
                content = m.group(1)
        p = text_frame.paragraphs[0] if first else text_frame.add_paragraph()
        first = False
        p.level = min(indent // 2, 8)
        for text, bold in _runs(content):
            r = p.add_run()
            r.text = text
            if bold or bold_all:
                r.font.bold = True


def _plain(text_frame):
    return [p.text for p in text_frame.paragraphs if p.text.strip()]


def _expected_fields(role, item, course):
    """field -> markdown text for a slide's role."""
    if role == "title":
        return {"title": course.title, "body": title_body(course)}
    if role == "section":
        return {"title": item.title}
    if role in ("concept", "concept-with-figure"):
        return {"title": item.title, "body": item.body, "citation": format_citation(item.sources), "notes": notes_text(item)}
    if role == "summary":
        return {"title": "Summary", "body": "\n".join("- " + c for c in course_concepts(course))}
    return {"title": "Sources", "body": "\n".join(format_deck_line(d) for d in decks(course))}


def _drop_template_slides(prs):
    ids = prs.slides._sldIdLst
    for sld in list(ids):
        prs.part.drop_rel(sld.rId)
        ids.remove(sld)


def export_deck(course: Course, mapping_path, out_path, resolve_image) -> Path:
    out_path = Path(out_path)
    mapping = load_mapping(mapping_path)
    problems = validate(mapping, course)
    if problems:
        raise DeckExportError(problems)

    images: Dict[str, bytes] = {}
    for leaf in iter_leaves(course):
        if not leaf.figure:
            continue
        try:
            data = resolve_image(leaf.figure.ref)
        except ImageRefError as e:
            problems.append(f"leaf '{leaf.title}': {e}")
            continue
        except Exception:
            data = None
        if not data:
            problems.append(f"image '{leaf.figure.ref}' on '{leaf.title}' could not be resolved")
        else:
            images[leaf.id] = data
    if problems:
        raise DeckExportError(problems)

    prs = Presentation(str(mapping.template_path))
    layouts = {name: found[0] for name, found in _layouts_by_name(prs).items()}
    _drop_template_slides(prs)
    for role, item in iter_slides(course):
        spec = mapping.roles[role]
        fields = dict(spec.get("placeholders") or {})
        slide = prs.slides.add_slide(layouts[spec["layout"]])
        values = _expected_fields(role, item, course)
        used = set()
        for f, idx in fields.items():
            if f == "image":
                slide.placeholders[idx].insert_picture(io.BytesIO(images[item.id]))
                used.add(idx)
            elif f == "notes":
                if idx == NOTES_PAGE:
                    slide.notes_slide.notes_text_frame.text = values["notes"]
                else:
                    write_text(slide.placeholders[idx].text_frame, values["notes"])
                    used.add(idx)
            elif f in values:
                write_text(slide.placeholders[idx].text_frame, values[f])
                used.add(idx)
        for ph in list(slide.placeholders):
            if ph.placeholder_format.idx not in used:
                ph._element.getparent().remove(ph._element)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    tmp = out_path.with_name(out_path.name + ".part")
    try:
        prs.save(str(tmp))
        problems = check_round_trip(course, mapping, tmp)
        if problems:
            raise DeckExportError(problems)
        os.replace(tmp, out_path)
    finally:
        if tmp.exists():
            tmp.unlink()
    return out_path


def _slide_text(slide, mapping, role, field):
    """The paragraphs a slide holds for a field, or None when the slide has no such text."""
    idx = (mapping.roles[role].get("placeholders") or {}).get(field)
    if field == "notes" and idx in (None, NOTES_PAGE):
        if not slide.has_notes_slide:
            return None
        return _plain(slide.notes_slide.notes_text_frame)
    for ph in slide.placeholders:
        if ph.placeholder_format.idx == idx and ph.has_text_frame:
            return _plain(ph.text_frame)
    return None


def _has_picture(slide):
    return any(sp.tag.endswith("}pic") for sp in slide.shapes._spTree)


def check_round_trip(course, mapping, pptx_path) -> List[str]:
    """Contract section 6: read the written deck back and compare it with the course."""
    if not isinstance(mapping, Mapping):
        mapping = load_mapping(mapping)
    problems: List[str] = []
    prs = Presentation(str(pptx_path))
    slides = list(prs.slides)
    expected = list(iter_slides(course))

    want = [mapping.roles[role]["layout"] for role, _ in expected]
    got = [s.slide_layout.name for s in slides]
    if want != got:
        problems.append(f"slide sequence: expected {len(want)} slides with layouts {want}, found {len(got)} with {got}")
        return problems

    seen: Dict[str, List[int]] = {}
    for n, (slide, (role, item)) in enumerate(zip(slides, expected), start=1):
        if role in ("concept", "concept-with-figure"):
            lines = _slide_text(slide, mapping, role, "notes") or []
            line = next((ln for ln in reversed(lines) if ln.startswith("Concepts:")), None)
            actual = [c.strip() for c in line[len("Concepts:"):].split(";") if c.strip()] if line else []
            if line is None:
                problems.append(f"slide {n} ('{item.title}'): no 'Concepts:' line in the notes")
            for c in item.concepts:
                if c not in actual:
                    problems.append(f"slide {n} ('{item.title}'): concept '{c}' is missing from the notes")
            for c in actual:
                if c not in item.concepts:
                    problems.append(f"slide {n} ('{item.title}'): concept '{c}' is not in the course for this leaf")
                seen.setdefault(c, []).append(n)

            cite = _slide_text(slide, mapping, role, "citation")
            want_cite = format_citation(item.sources)
            if cite != [want_cite]:
                problems.append(f"slide {n} ('{item.title}'): citation is {cite!r}, expected {want_cite!r}")

        has_pic = _has_picture(slide)
        if role == "concept-with-figure" and not has_pic:
            problems.append(f"slide {n} ('{item.title}'): the figure is missing")
        if role != "concept-with-figure" and has_pic:
            problems.append(f"slide {n} ({role}): holds a picture but should not")

        if role == "summary":
            body = _slide_text(slide, mapping, role, "body")
            if body != course_concepts(course):
                problems.append(f"slide {n}: the summary lists {body}, the course has {course_concepts(course)}")
        if role == "sources":
            body = _slide_text(slide, mapping, role, "body")
            want_lines = [format_deck_line(d) for d in decks(course)]
            if body != want_lines:
                problems.append(f"slide {n}: the sources slide lists {body}, the course cites {want_lines}")

    for c, where in seen.items():
        if len(where) > 1:
            problems.append(f"concept '{c}' appears on {len(where)} slides: {where}")
    for c in course_concepts(course):
        if c not in seen:
            problems.append(f"concept '{c}' appears on no concept slide")
    return problems
