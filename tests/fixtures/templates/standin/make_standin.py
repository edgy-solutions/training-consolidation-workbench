"""Builds standin.pptx: a plain six-layout template for the deck-template contract.

Run: uv run python tests/fixtures/templates/standin/make_standin.py [out.pptx]
"""
import copy
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

LAYOUT_NAMES = ["title", "section", "concept", "concept-with-figure", "summary", "sources"]
DROP_IDX = {10, 11, 12}  # date, footer, slide number

# default-template layout index -> new name
SOURCE_LAYOUT = {0: "title", 5: "section", 1: "concept", 3: "concept-with-figure", 7: "summary", 9: "sources"}

W = 9144000
FULL = (457200, 274638, 8229600, 1143000)
BODY = (457200, 1600200, 8229600, 4300000)
CITE = (457200, 6000000, 8229600, 500000)


def _ph(shape):
    return shape._element.nvSpPr.nvPr.ph


def _place(shape, box):
    shape.left, shape.top, shape.width, shape.height = (Emu(v) for v in box)


def _set_type(shape, ph_type, idx=None):
    ph = _ph(shape)
    if "type" in ph.attrib:
        del ph.attrib["type"]
    if ph_type:
        ph.set("type", ph_type)
    if idx is not None:
        ph.set("idx", str(idx))


def _by_idx(layout, idx):
    for s in layout.placeholders:
        if s.placeholder_format.idx == idx:
            return s
    raise KeyError(idx)


def _strip(layout, keep):
    for s in list(layout.placeholders):
        if s.placeholder_format.idx not in keep:
            s._element.getparent().remove(s._element)


def build(out_path):
    prs = Presentation()
    layouts = list(prs.slide_layouts)
    for i, layout in enumerate(layouts):
        if i not in SOURCE_LAYOUT:
            prs.slide_layouts.remove(layout)
    by_name = {}
    for i, name in SOURCE_LAYOUT.items():
        layout = layouts[i]
        layout._element.cSld.set("name", name)
        by_name[name] = layout

    lay = by_name["title"]
    _strip(lay, {0, 1})

    lay = by_name["section"]
    _strip(lay, {0})

    lay = by_name["concept"]
    _strip(lay, {0, 1})
    body = _by_idx(lay, 1)
    _place(body, (457200, 1600200, 8229600, 4300000))
    cite = copy.deepcopy(body._element)
    lay.shapes._spTree.append(cite)
    cite.nvSpPr.cNvPr.set("id", "90")
    cite.nvSpPr.cNvPr.set("name", "Citation Placeholder")
    ph = cite.nvSpPr.nvPr.ph
    ph.set("idx", "2")
    ph.set("type", "body")
    _place(_by_idx(lay, 2), CITE)

    lay = by_name["concept-with-figure"]
    _strip(lay, {0, 1, 2})
    _place(_by_idx(lay, 1), (457200, 1600200, 4038600, 4300000))
    cite = _by_idx(lay, 2)
    _set_type(cite, "body")
    _place(cite, CITE)
    pic = copy.deepcopy(_by_idx(lay, 1)._element)
    lay.shapes._spTree.append(pic)
    pic.nvSpPr.cNvPr.set("id", "91")
    pic.nvSpPr.cNvPr.set("name", "Picture Placeholder")
    pic.nvSpPr.nvPr.ph.set("idx", "3")
    pic.nvSpPr.nvPr.ph.set("type", "pic")
    if pic.txBody is not None:
        pic.remove(pic.txBody)
    _place(_by_idx(lay, 3), (4648200, 1600200, 4038600, 4300000))

    lay = by_name["summary"]
    _strip(lay, {0, 1})
    _place(_by_idx(lay, 0), FULL)
    _place(_by_idx(lay, 1), BODY)

    lay = by_name["sources"]
    _strip(lay, {0, 1})
    for s in lay.placeholders:
        bodyPr = s._element.txBody.bodyPr
        if "vert" in bodyPr.attrib:
            del bodyPr.attrib["vert"]

    out = Path(out_path)
    prs.save(str(out))
    return out


if __name__ == "__main__":
    target = sys.argv[1] if len(sys.argv) > 1 else Path(__file__).with_name("standin.pptx")
    build(target)
