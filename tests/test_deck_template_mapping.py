import pytest
import yaml
from pptx import Presentation

from src.publishing.course import DeckExportError
from src.publishing.deck_export import export_deck
from tests.deck_fixtures import copy_standin, golden_course, resolver


def _broken(tmp_path, edit):
    path = copy_standin(tmp_path / "tpl")
    data = yaml.safe_load(path.read_text(encoding="utf-8"))
    edit(data)
    path.write_text(yaml.safe_dump(data), encoding="utf-8")
    return path


def _fails(tmp_path, mapping_path):
    out = tmp_path / "out.pptx"
    with pytest.raises(DeckExportError) as e:
        export_deck(golden_course(), mapping_path, out, resolver())
    assert not out.exists()
    assert not list(tmp_path.glob("*.part"))
    return e.value.problems


def test_unmapped_role(tmp_path):
    problems = _fails(tmp_path, _broken(tmp_path, lambda d: d["roles"].pop("concept-with-figure")))
    assert "role 'concept-with-figure' is not mapped" in problems


def test_unknown_role(tmp_path):
    problems = _fails(tmp_path, _broken(tmp_path, lambda d: d["roles"].update(
        bogus={"layout": "title", "placeholders": {"title": 0}})))
    assert any("unknown role 'bogus'" in p for p in problems)


def test_unknown_field(tmp_path):
    problems = _fails(tmp_path, _broken(tmp_path, lambda d: d["roles"]["section"]["placeholders"].update(subtitle=1)))
    assert any("role 'section'" in p and "unknown field 'subtitle'" in p for p in problems)


def test_missing_layout_lists_the_templates_layouts(tmp_path):
    problems = _fails(tmp_path, _broken(tmp_path, lambda d: d["roles"]["summary"].update(layout="nope")))
    msg = next(p for p in problems if "'nope'" in p)
    for name in ("title", "section", "concept", "concept-with-figure", "summary", "sources"):
        assert name in msg


def test_ambiguous_layout(tmp_path):
    path = _broken(tmp_path, lambda d: None)
    tpl = path.parent / "standin.pptx"
    prs = Presentation(str(tpl))
    prs.slide_layouts.get_by_name("section")._element.cSld.set("name", "title")
    prs.save(str(tpl))
    problems = _fails(tmp_path, path)
    assert any("layout 'title' is carried by 2 layouts" in p for p in problems)


def test_bad_placeholder_idx(tmp_path):
    problems = _fails(tmp_path, _broken(tmp_path, lambda d: d["roles"]["concept"]["placeholders"].update(body=9)))
    assert any("role 'concept'" in p and "placeholder 9" in p for p in problems)


def test_image_on_a_placeholder_that_cannot_take_a_picture(tmp_path):
    problems = _fails(tmp_path, _broken(tmp_path, lambda d: d["roles"]["concept-with-figure"]["placeholders"].update(image=2)))
    assert any("cannot take a picture" in p for p in problems)


def test_required_field_unmapped(tmp_path):
    problems = _fails(tmp_path, _broken(tmp_path, lambda d: d["roles"]["concept"]["placeholders"].pop("citation")))
    assert any("role 'concept'" in p and "required field 'citation'" in p for p in problems)


def test_content_with_no_field_to_go_to(tmp_path):
    problems = _fails(tmp_path, _broken(tmp_path, lambda d: d["roles"]["title"]["placeholders"].pop("body")))
    assert any("role 'title'" in p and "subtitle" in p for p in problems)


def test_wrong_version(tmp_path):
    problems = _fails(tmp_path, _broken(tmp_path, lambda d: d.update(deck_template=2)))
    assert any("deck_template" in p for p in problems)


def test_missing_version(tmp_path):
    problems = _fails(tmp_path, _broken(tmp_path, lambda d: d.pop("deck_template")))
    assert any("deck_template" in p for p in problems)


def test_missing_template_file(tmp_path):
    problems = _fails(tmp_path, _broken(tmp_path, lambda d: d.update(template="nope.pptx")))
    assert any("nope.pptx" in p and "missing" in p for p in problems)


def test_all_problems_are_reported_together(tmp_path):
    def edit(d):
        d["roles"].pop("summary")
        d["roles"]["concept"]["placeholders"].pop("citation")
        d["roles"]["section"]["layout"] = "nope"
    problems = _fails(tmp_path, _broken(tmp_path, edit))
    assert len(problems) >= 3


def test_unresolvable_image_writes_nothing(tmp_path):
    from tests.deck_fixtures import copy_standin as cp
    out = tmp_path / "out.pptx"
    with pytest.raises(DeckExportError) as e:
        export_deck(golden_course(), cp(tmp_path / "t"), out, lambda ref: None)
    assert e.value.problems == ["image 'fixture://images/slide_4_img_0.png' on 'Reservoir Schematic' could not be resolved"]
    assert not out.exists()
