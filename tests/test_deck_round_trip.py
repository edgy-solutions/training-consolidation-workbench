import copy

import pytest
from pptx import Presentation

from src.publishing.deck_export import DEFAULT_MAPPING, check_round_trip, export_deck, load_mapping
from tests.deck_fixtures import golden_course, resolver

MAPPING = load_mapping(DEFAULT_MAPPING)


@pytest.fixture(scope="module")
def exported(tmp_path_factory):
    course = golden_course()
    out = export_deck(course, DEFAULT_MAPPING, tmp_path_factory.mktemp("rt") / "g.pptx", resolver())
    return course, out


def _tamper(exported, tmp_path, fn):
    course, out = exported
    prs = Presentation(str(out))
    fn(prs)
    path = tmp_path / "t.pptx"
    prs.save(str(path))
    return check_round_trip(course, MAPPING, path)


def _ph(slide, idx):
    return next(p for p in slide.placeholders if p.placeholder_format.idx == idx)


def test_untouched_is_clean(exported, tmp_path):
    assert _tamper(exported, tmp_path, lambda prs: None) == []


def test_deleted_citation(exported, tmp_path):
    def fn(prs):
        _ph(prs.slides[2], 2).text_frame.text = ""
    assert any("citation" in p for p in _tamper(exported, tmp_path, fn))


def test_extra_concept_in_notes(exported, tmp_path):
    def fn(prs):
        prs.slides[3].notes_slide.notes_text_frame.text += "; Invented concept"
    assert any("Invented concept" in p for p in _tamper(exported, tmp_path, fn))


def test_concept_repeated_on_two_slides(exported, tmp_path):
    def fn(prs):
        prs.slides[3].notes_slide.notes_text_frame.text += "; Gear pump; Hydraulic pressure"
    problems = _tamper(exported, tmp_path, fn)
    assert any("'Hydraulic pressure' appears on 2 slides" in p for p in problems)


def test_picture_removed(exported, tmp_path):
    def fn(prs):
        tree = prs.slides[5].shapes._spTree
        for sp in [s for s in tree if s.tag.endswith("}pic")]:
            tree.remove(sp)
    assert any("figure is missing" in p for p in _tamper(exported, tmp_path, fn))


def test_picture_added_to_concept_slide(exported, tmp_path):
    def fn(prs):
        pic = next(s for s in prs.slides[5].shapes._spTree if s.tag.endswith("}pic"))
        prs.slides[2].shapes._spTree.append(copy.deepcopy(pic))
    assert any("holds a picture" in p for p in _tamper(exported, tmp_path, fn))


def test_extra_slide(exported, tmp_path):
    def fn(prs):
        prs.slides.add_slide(prs.slide_layouts.get_by_name("section"))
    assert any("slide sequence" in p for p in _tamper(exported, tmp_path, fn))


def test_layout_order_changed(exported, tmp_path):
    def fn(prs):
        ids = prs.slides._sldIdLst
        items = list(ids)
        ids.remove(items[1])
        ids.insert(3, items[1])
    assert any("slide sequence" in p for p in _tamper(exported, tmp_path, fn))
