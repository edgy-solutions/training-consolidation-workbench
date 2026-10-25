import importlib.util

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER as T

from tests.deck_fixtures import STANDIN_DIR

EXPECTED = {
    "title": {0: T.CENTER_TITLE, 1: T.SUBTITLE},
    "section": {0: T.TITLE},
    "concept": {0: T.TITLE, 1: T.OBJECT, 2: T.BODY},
    "concept-with-figure": {0: T.TITLE, 1: T.OBJECT, 2: T.BODY, 3: T.PICTURE},
    "summary": {0: T.TITLE, 1: T.OBJECT},
    "sources": {0: T.TITLE, 1: T.BODY},
}


def _shape(path):
    prs = Presentation(str(path))
    assert len(prs.slides) == 0
    assert len(prs.slide_masters) == 1
    return {l.name: {p.placeholder_format.idx: p.placeholder_format.type for p in l.placeholders}
            for l in prs.slide_layouts}


def test_committed_standin_has_the_six_layouts():
    assert _shape(STANDIN_DIR / "standin.pptx") == EXPECTED


def test_regenerating_gives_the_same_layouts(tmp_path):
    spec = importlib.util.spec_from_file_location("make_standin", STANDIN_DIR / "make_standin.py")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    out = tmp_path / "standin.pptx"
    mod.build(out)
    assert _shape(out) == EXPECTED
