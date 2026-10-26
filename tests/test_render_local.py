"""Render without Dagster: the in-process runner, its job store and the routes."""
import json
import re
import shutil
import subprocess
import sys
import time
from io import BytesIO
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from pptx import Presentation

from src.publishing import render_jobs
from src.publishing.render import COURSE_NODES_QUERY, render_project
from src.publishing.render_jobs import RUN_ID_RE, create_job
from tests.deck_fixtures import FIXTURE_DIR, course_data, reader

LAYOUTS = ["title", "section", "concept", "concept", "section", "concept-with-figure", "summary", "sources"]
PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
REPO = Path(__file__).resolve().parent.parent


class FakeNeo4j:
    def __init__(self):
        data = course_data()
        self.project = data["project"]
        self.nodes = data["nodes"]

    def execute_query(self, query, params=None):
        if "-[:HAS_CHILD*0..]->(n)" in query:
            return [{"n.id": self.project["id"]}]
        if "RETURN p.title AS title" in query:
            if params["id"] != self.project["id"]:
                return []
            return [{"title": self.project["title"], "deck_ids": self.project["deck_ids"]}]
        if query == COURSE_NODES_QUERY:
            return [dict(n) for n in self.nodes]
        return []


def pptx_layouts(prs):
    return [s.slide_layout.name for s in prs.slides]


@pytest.fixture
def harness(monkeypatch, tmp_path):
    # The default fixture dir is a multi-course reader that prefixes asset names; the golden
    # course alone (the root files, no sub-fixtures) resolves its images by bare name.
    single = tmp_path / "fixture"
    single.mkdir()
    for f in FIXTURE_DIR.iterdir():
        if f.is_file():
            shutil.copy(f, single / f.name)
        elif f.name in ("images", "assets"):
            shutil.copytree(f, single / f.name)
    monkeypatch.setenv("SOURCE_READER", "fixture")
    monkeypatch.setenv("SOURCE_FIXTURE_DIR", str(single))
    exports = tmp_path / "exports"
    exports.mkdir()
    monkeypatch.setenv("WB_EXPORT_DIR", str(exports))
    monkeypatch.setenv("WB_STORE_WRITES", "on")
    monkeypatch.delenv("DAGSTER_HOST", raising=False)
    monkeypatch.delenv("RENDER_RUNNER", raising=False)
    from src.auth.security import User, get_current_user
    from src.workbench import main

    fake = FakeNeo4j()
    monkeypatch.setattr(main.neo4j_client, "execute_query", fake.execute_query, raising=False)
    monkeypatch.setattr(main, "RENDER_POLL_SECONDS", 0.01)
    main.app.dependency_overrides[get_current_user] = lambda: User(id="u1", email="u1@example.com")
    yield TestClient(main.app, raise_server_exceptions=False), fake, main
    main.app.dependency_overrides.pop(get_current_user, None)


def events(text):
    """SSE body -> list of (event, data dict)."""
    out = []
    for block in re.split(r"\r?\n\r?\n", text.strip()):
        ev, data = None, None
        for line in block.splitlines():
            if line.startswith("event:"):
                ev = line.split(":", 1)[1].strip()
            elif line.startswith("data:"):
                data = json.loads(line.split(":", 1)[1].strip())
        if ev:
            out.append((ev, data))
    return out


def trigger(client, **extra):
    return client.post("/render/trigger", json={"project_id": "proj-hydraulics", **extra})


def test_render_project_writes_the_golden_deck(monkeypatch, tmp_path):
    monkeypatch.setenv("SOURCE_READER", "fixture")
    out = tmp_path / "deck.pptx"
    render_project("proj-hydraulics", "pptx", None, out, neo4j_client=FakeNeo4j(), minio_client=None, reader=reader())
    prs = Presentation(str(out))
    assert len(prs.slides) == 8
    assert pptx_layouts(prs) == LAYOUTS


def test_render_project_missing_project_and_bad_format(tmp_path):
    with pytest.raises(LookupError, match="Project nope not found"):
        render_project("nope", "pptx", None, tmp_path / "x", neo4j_client=FakeNeo4j(), minio_client=None, reader=reader())
    with pytest.raises(ValueError):
        render_project("proj-hydraulics", "pdf", None, tmp_path / "x", neo4j_client=FakeNeo4j(), minio_client=None, reader=reader())


def test_trigger_events_and_file_end_to_end(harness):
    client, _, _ = harness
    r = trigger(client)
    assert r.status_code == 200
    body = r.json()
    assert body["job"] == "render_local"
    run_id = body["run_id"]
    assert RUN_ID_RE.match(run_id)
    assert body["filename"].startswith("Hydraulic_Systems__consolidated__")
    assert "Untitled_Project" not in body["filename"]

    evs = events(client.get(f"/render/events/{run_id}").text)
    complete = [d for e, d in evs if e == "complete"]
    assert len(complete) == 1, evs
    assert complete[0]["download_url"] == f"/render/file/{run_id}"

    f = client.get(complete[0]["download_url"])
    assert f.status_code == 200
    assert f.headers["content-type"].startswith(PPTX)
    assert len(Presentation(BytesIO(f.content)).slides) == 8


def test_failure_reports_the_leaf_and_leaves_no_file(harness, tmp_path):
    tmp_path = tmp_path / "exports"
    client, fake, _ = harness
    leaf = next(n for n in fake.nodes if n["content_markdown"] and n["key_concepts"])
    fake.nodes = [{**n, "content_markdown": ""} if n is leaf else n for n in fake.nodes]
    run_id = trigger(client).json()["run_id"]

    evs = events(client.get(f"/render/events/{run_id}").text)
    errors = [d for e, d in evs if e == "error"]
    assert len(errors) == 1
    assert errors[0]["status"] == "FAILURE"
    assert any(leaf["title"] in p for p in errors[0]["problems"])
    assert client.get(f"/render/file/{run_id}").status_code == 404
    assert not list((tmp_path / run_id).glob("*.pptx"))


def test_file_route_rejects_bad_ids(harness):
    client, _, _ = harness
    assert client.get("/render/file/../x").status_code == 404
    assert client.get("/render/file/local-zz").status_code == 404
    assert client.get("/render/file/local-" + "0" * 32).status_code == 404


def test_events_for_an_unknown_local_id_is_404(harness):
    client, _, _ = harness
    assert client.get("/render/events/local-" + "0" * 32).status_code == 404
    assert client.get("/render/events/local-zz").status_code == 404


class StubDagster:
    def __init__(self):
        self.calls = []

    def execute_query(self, query, variables=None):
        self.calls.append("execute_query")
        return {"addDynamicPartition": {"__typename": "AddDynamicPartitionSuccess"}}

    def submit_job_execution(self, *a, **k):
        self.calls.append("submit_job_execution")
        return "dagster-run-1"


def test_runner_choice(harness, monkeypatch, tmp_path):
    tmp_path = tmp_path / "exports"
    client, _, main = harness
    stub = StubDagster()
    monkeypatch.setattr(main, "dagster_client", stub)

    monkeypatch.setenv("DAGSTER_HOST", "dagster.example")
    r = trigger(client)
    assert r.status_code == 200
    assert r.json()["job"] == "render_asset_job"
    assert "submit_job_execution" in stub.calls
    assert not [p for p in tmp_path.iterdir() if p.name.startswith("local-")]

    monkeypatch.setenv("RENDER_RUNNER", "local")
    assert trigger(client).json()["job"] == "render_local"

    monkeypatch.setenv("RENDER_RUNNER", "bogus")
    r = trigger(client)
    assert r.status_code == 500
    assert r.json()["detail"] == "RENDER_RUNNER must be local or dagster"


def test_unknown_format_is_400(harness):
    client, _, _ = harness
    r = trigger(client, format="pdf")
    assert r.status_code == 400
    assert r.json()["detail"] == "format must be one of pptx, typ"


def test_render_modules_do_not_import_dagster():
    code = ("import sys, src.publishing.render, src.publishing.render_jobs; "
            "sys.exit(any(m == 'dagster' or m.startswith('dagster.') for m in sys.modules))")
    assert subprocess.run([sys.executable, "-c", code], cwd=str(REPO)).returncode == 0


def test_create_job_sweeps_old_jobs(monkeypatch, tmp_path):
    monkeypatch.setenv("WB_EXPORT_DIR", str(tmp_path))
    old = create_job("p", "a.pptx", "pptx", "u1")["run_id"]
    fresh = create_job("p", "b.pptx", "pptx", "u1")["run_id"]
    rec = json.loads((tmp_path / old / "job.json").read_text())
    rec["created"] = time.time() - 25 * 3600
    (tmp_path / old / "job.json").write_text(json.dumps(rec))
    create_job("p", "c.pptx", "pptx", "u1")
    assert not (tmp_path / old).exists()
    assert (tmp_path / fresh).exists()
    assert render_jobs.read_job(fresh)["status"] == "QUEUED"
