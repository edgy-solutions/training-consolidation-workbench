from pptx import Presentation

from src.publishing.deck_export import DEFAULT_MAPPING, check_round_trip, export_deck, load_mapping
from tests.deck_fixtures import golden_course, resolver


def _ph(slide, idx):
    return next(p for p in slide.placeholders if p.placeholder_format.idx == idx)


def _export(tmp_path):
    course = golden_course()
    out = export_deck(course, DEFAULT_MAPPING, tmp_path / "out.pptx", resolver())
    return course, Presentation(str(out)), out


def test_golden_export_has_the_eight_slides(tmp_path):
    course, prs, out = _export(tmp_path)
    slides = list(prs.slides)
    assert [s.slide_layout.name for s in slides] == [
        "title", "section", "concept", "concept", "section", "concept-with-figure", "summary", "sources"]
    assert _ph(slides[0], 0).text_frame.text == "Hydraulic Systems (consolidated)"
    assert _ph(slides[0], 1).text_frame.text == "Consolidated from 1 source deck"
    assert _ph(slides[1], 0).text_frame.text == "Hydraulic Principles"
    assert _ph(slides[2], 0).text_frame.text == "Pascal's Law"
    assert _ph(slides[2], 2).text_frame.text == "Sources: hydraulics_fixture.pptx, slides 1, 2"
    assert _ph(slides[3], 2).text_frame.text == "Source: hydraulics_fixture.pptx, slide 3"
    assert _ph(slides[4], 0).text_frame.text == "System Components"
    assert _ph(slides[5], 0).text_frame.text == "Reservoir Schematic"
    assert _ph(slides[5], 2).text_frame.text == "Source: hydraulics_fixture.pptx, slide 4"
    assert _ph(slides[5], 1).text_frame.text == "Figure 1: reservoir, filter and pump in series"
    assert [p.text for p in _ph(slides[6], 1).text_frame.paragraphs] == [
        "Pascal's law", "Hydraulic pressure", "Hydraulic pump", "Gear pump", "Hydraulic reservoir"]
    assert _ph(slides[7], 1).text_frame.text == "hydraulics_fixture.pptx: slides 1, 2, 3, 4"
    assert "Stray" not in "".join(p.text_frame.text for s in slides for p in s.placeholders if p.has_text_frame)


def test_golden_pictures_notes_and_round_trip(tmp_path):
    course, prs, out = _export(tmp_path)
    slides = list(prs.slides)
    has_pic = [any(sp.tag.endswith("}pic") for sp in s.shapes._spTree) for s in slides]
    assert has_pic == [False] * 5 + [True] + [False] * 2
    assert slides[2].notes_slide.notes_text_frame.text.splitlines()[-1] == "Concepts: Pascal's law; Hydraulic pressure"
    assert slides[3].notes_slide.notes_text_frame.text.splitlines()[-1] == "Concepts: Hydraulic pump; Gear pump"
    assert slides[5].notes_slide.notes_text_frame.text.splitlines()[-1] == "Concepts: Hydraulic reservoir"
    assert check_round_trip(course, load_mapping(DEFAULT_MAPPING), out) == []


def test_no_part_file_left_behind(tmp_path):
    _export(tmp_path)
    assert sorted(p.name for p in tmp_path.iterdir()) == ["out.pptx"]
